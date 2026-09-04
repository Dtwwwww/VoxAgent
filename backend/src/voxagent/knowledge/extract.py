from __future__ import annotations

import zipfile
from dataclasses import dataclass
from pathlib import Path

MAX_DOCUMENT_BYTES = 20 * 1024 * 1024

_MIME_TYPES = {
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".markdown": "text/markdown",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


@dataclass(frozen=True, slots=True)
class DocumentSection:
    content: str
    page_number: int | None = None


@dataclass(frozen=True, slots=True)
class ExtractedDocument:
    display_name: str
    mime_type: str
    sections: tuple[DocumentSection, ...]


@dataclass(frozen=True, slots=True)
class DocumentChunk:
    ordinal: int
    content: str
    page_number: int | None


def extract_document(path: Path) -> ExtractedDocument:
    source = path.expanduser().resolve(strict=True)
    if not source.is_file():
        raise ValueError("所选路径不是文件")
    if source.stat().st_size > MAX_DOCUMENT_BYTES:
        raise ValueError("文件不能超过 20 MB")
    suffix = source.suffix.lower()
    mime_type = _MIME_TYPES.get(suffix)
    if mime_type is None:
        raise ValueError("不支持该文件格式，仅支持 TXT、Markdown、PDF 和 DOCX")

    if suffix in {".txt", ".md", ".markdown"}:
        try:
            content = source.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError as error:
            raise ValueError("文本文件必须使用 UTF-8 编码") from error
        sections = (DocumentSection(content.replace("\r\n", "\n").strip()),)
    elif suffix == ".pdf":
        sections = _extract_pdf(source)
    else:
        sections = _extract_docx(source)

    if not any(section.content.strip() for section in sections):
        raise ValueError("文档中未检测到可提取文字，暂不支持图片 OCR")
    return ExtractedDocument(source.name, mime_type, sections)


def _extract_pdf(path: Path) -> tuple[DocumentSection, ...]:
    if path.read_bytes()[:5] != b"%PDF-":
        raise ValueError("PDF 文件内容与扩展名不匹配")
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(path)
        sections = tuple(
            DocumentSection((page.extract_text() or "").strip(), page_number=index)
            for index, page in enumerate(reader.pages, start=1)
        )
    except (OSError, PdfReadError, ValueError) as error:
        raise ValueError("PDF 文件无法解析") from error
    return sections


def _extract_docx(path: Path) -> tuple[DocumentSection, ...]:
    if not zipfile.is_zipfile(path):
        raise ValueError("DOCX 文件内容与扩展名不匹配")
    from docx import Document
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = Document(path)
    except (OSError, ValueError, PackageNotFoundError) as error:
        raise ValueError("DOCX 文件无法解析") from error
    paragraphs = tuple(item.text.strip() for item in document.paragraphs if item.text.strip())
    return (DocumentSection("\n\n".join(paragraphs)),)


def chunk_document(
    document: ExtractedDocument,
    target_chars: int = 600,
    overlap_chars: int = 80,
) -> tuple[DocumentChunk, ...]:
    if target_chars < 1:
        raise ValueError("target_chars must be positive")
    if overlap_chars < 0 or overlap_chars >= target_chars:
        raise ValueError("overlap_chars must be non-negative and smaller than target_chars")
    step = target_chars - overlap_chars
    chunks: list[DocumentChunk] = []
    for section in document.sections:
        content = section.content
        for start in range(0, len(content), step):
            value = content[start : start + target_chars]
            if value.strip():
                chunks.append(DocumentChunk(len(chunks), value, section.page_number))
            if start + target_chars >= len(content):
                break
    return tuple(chunks)
