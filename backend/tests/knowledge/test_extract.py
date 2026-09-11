from __future__ import annotations

from pathlib import Path

import pytest

from voxagent.knowledge.extract import (
    DocumentSection,
    ExtractedDocument,
    chunk_document,
    extract_document,
)


def _pdf_bytes(page_texts: tuple[str, ...]) -> bytes:
    objects: list[bytes] = []
    page_ids = [4 + index * 2 for index in range(len(page_texts))]
    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>".encode())
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for page_id, text in zip(page_ids, page_texts, strict=True):
        content_id = page_id + 1
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("ascii")
        objects.append(
            (
                f"<< /Type /Page /Parent 2 0 R /Resources << /Font << /F1 3 0 R >> >> "
                f"/MediaBox [0 0 612 792] /Contents {content_id} 0 R >>"
            ).encode()
        )
        objects.append(
            b"<< /Length "
            + str(len(stream)).encode()
            + b" >>\nstream\n"
            + stream
            + b"\nendstream"
        )

    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for object_id, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{object_id} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(
        f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return bytes(output)


def test_extracts_utf8_text_and_markdown_headings(tmp_path: Path) -> None:
    text_path = tmp_path / "资料.txt"
    text_path.write_text("第一段\n\n第二段", encoding="utf-8")
    markdown_path = tmp_path / "说明.md"
    markdown_path.write_text("# 标题\n\n正文内容", encoding="utf-8")

    text = extract_document(text_path)
    markdown = extract_document(markdown_path)

    assert text.mime_type == "text/plain"
    assert "第一段\n\n第二段" in text.sections[0].content
    assert markdown.mime_type == "text/markdown"
    assert markdown.sections[0].content.startswith("# 标题")


def test_extracts_two_page_text_pdf_with_page_numbers(tmp_path: Path) -> None:
    path = tmp_path / "two-pages.pdf"
    path.write_bytes(_pdf_bytes(("first page", "second page")))

    document = extract_document(path)

    assert [(section.page_number, section.content) for section in document.sections] == [
        (1, "first page"),
        (2, "second page"),
    ]


def test_extracts_docx_paragraphs(tmp_path: Path) -> None:
    from docx import Document

    path = tmp_path / "notes.docx"
    source = Document()
    source.add_paragraph("第一段")
    source.add_paragraph("第二段")
    source.save(path)

    document = extract_document(path)

    expected_mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert document.mime_type == expected_mime
    assert document.sections[0].content == "第一段\n\n第二段"


def test_rejects_disallowed_extensions_and_mime_mismatches(tmp_path: Path) -> None:
    executable = tmp_path / "notes.exe"
    executable.write_text("not allowed", encoding="utf-8")
    fake_pdf = tmp_path / "fake.pdf"
    fake_pdf.write_text("not a pdf", encoding="utf-8")

    with pytest.raises(ValueError, match="不支持"):
        extract_document(executable)
    with pytest.raises(ValueError, match="PDF"):
        extract_document(fake_pdf)


def test_rejects_files_larger_than_twenty_megabytes(tmp_path: Path) -> None:
    path = tmp_path / "too-large.txt"
    with path.open("wb") as stream:
        stream.seek(20 * 1024 * 1024)
        stream.write(b"x")

    with pytest.raises(ValueError, match="20 MB"):
        extract_document(path)


def test_rejects_image_only_pdf_with_chinese_explanation(tmp_path: Path) -> None:
    from pypdf import PdfWriter

    path = tmp_path / "image-only.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    with path.open("wb") as stream:
        writer.write(stream)

    with pytest.raises(ValueError, match="未检测到可提取文字"):
        extract_document(path)


def test_chunking_is_deterministic_and_preserves_page_and_heading() -> None:
    document = ExtractedDocument(
        display_name="guide.md",
        mime_type="text/markdown",
        sections=(DocumentSection("# 标题\n\nabcdefghijklmnop", page_number=3),),
    )

    first = chunk_document(document, target_chars=12, overlap_chars=3)
    second = chunk_document(document, target_chars=12, overlap_chars=3)

    assert first == second
    assert first[0].content.startswith("# 标题")
    assert all(chunk.page_number == 3 for chunk in first)
    assert first[0].content[-3:] == first[1].content[:3]
